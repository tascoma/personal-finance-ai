"""Tests for the orchestrate-parse phase — service + HTTP route.

The orchestrator LLM call is stubbed via monkeypatch on `run_orchestrator`;
sub-agents (statement/paystub/mortgage) and the classifier are stubbed where
needed so tests run without network or real Anthropic calls.
"""

import uuid
from datetime import date
from decimal import Decimal
from unittest.mock import AsyncMock

import openpyxl
import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.agents._base import AgentError
from app.agents.orchestrator import DocumentPlan, OrchestrationPlan
from app.core.config import settings
from app.databases import Base
from app.dependencies import get_current_user, get_db_session
from app.main import app
from app.models.account import Account
from app.models.document import Document
from app.models.raw_transaction import RawTransaction
from app.models.user import User
from app.services import classify as classify_service
from app.services import document as document_service
from app.services import orchestrate as orchestrate_service
from app.services import period as period_service

TEST_DB_URL = "sqlite+aiosqlite:///:memory:"


# ── fixtures ─────────────────────────────────────────────────────────────────


@pytest_asyncio.fixture
async def session_factory():
    eng = create_async_engine(TEST_DB_URL, echo=False)
    async with eng.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(eng, expire_on_commit=False)
    async with factory() as session:
        session.add_all(
            [
                Account(
                    account_code=100101,
                    account_name="Checking",
                    account_type="Asset",
                    sub_category="Cash",
                    normal_balance="debit",
                    is_memo=False,
                    is_active=True,
                ),
                Account(
                    account_code=200101,
                    account_name="Mastercard",
                    account_type="Liability",
                    sub_category="Credit Cards",
                    normal_balance="credit",
                    is_memo=False,
                    is_active=True,
                ),
            ]
        )
        await session.commit()
    yield factory
    await eng.dispose()


@pytest_asyncio.fixture
async def open_period(session_factory):
    async with session_factory() as session:
        period = await period_service.create_period(session, 2026, 1)
    return period


@pytest.fixture
def upload_root(tmp_path, monkeypatch):
    root = tmp_path / "uploads"
    monkeypatch.setattr(document_service, "UPLOAD_ROOT", root)
    return root


@pytest_asyncio.fixture
async def csv_bank_doc(session_factory, open_period, upload_root):
    period_dir = upload_root / str(open_period.period_id)
    period_dir.mkdir(parents=True)
    file_path = period_dir / "bank.csv"
    file_path.write_text(
        "Date,Description,ChkRef,Amount,Balance\n"
        "1/2/26,COFFEE SHOP,,($5.00),$995.00\n"
        '1/3/26,DIRECT DEPOSIT,,"$2,000.00","$2,995.00"\n',
    )
    async with session_factory() as session:
        doc = Document(
            period_id=open_period.period_id,
            document_type="bank_statement",
            file_name="bank.csv",
            file_path=str(file_path),
            source_account_code=100101,
            parse_status="pending",
        )
        session.add(doc)
        await session.commit()
        await session.refresh(doc)
    return doc


@pytest_asyncio.fixture
async def xlsx_card_doc(session_factory, open_period, upload_root):
    period_dir = upload_root / str(open_period.period_id)
    period_dir.mkdir(parents=True, exist_ok=True)
    file_path = period_dir / "card.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Date", "Transaction", "Name", "Memo", "Amount"])
    ws.append([date(2026, 1, 5), "DEBIT", "WALMART", "", -42.10])
    wb.save(file_path)

    async with session_factory() as session:
        doc = Document(
            period_id=open_period.period_id,
            document_type="credit_card",
            file_name="card.xlsx",
            file_path=str(file_path),
            source_account_code=200101,
            parse_status="pending",
        )
        session.add(doc)
        await session.commit()
        await session.refresh(doc)
    return doc


@pytest_asyncio.fixture
async def mislabeled_csv_doc(session_factory, open_period, upload_root):
    """A CSV uploaded with the wrong declared type (credit_card) — actually a bank statement."""
    period_dir = upload_root / str(open_period.period_id)
    period_dir.mkdir(parents=True, exist_ok=True)
    file_path = period_dir / "mislabeled.csv"
    file_path.write_text(
        "Date,Description,ChkRef,Amount,Balance\n"
        "1/4/26,RENT CHECK,,($1500.00),$1000.00\n",
    )
    async with session_factory() as session:
        doc = Document(
            period_id=open_period.period_id,
            document_type="credit_card",  # WRONG — orchestrator should correct this
            file_name="mislabeled.csv",
            file_path=str(file_path),
            source_account_code=100101,
            parse_status="pending",
        )
        session.add(doc)
        await session.commit()
        await session.refresh(doc)
    return doc


@pytest_asyncio.fixture
async def client(session_factory):
    async def override_get_db_session():
        async with session_factory() as session:
            yield session

    async def _mock_user() -> User:
        return User(user_id=uuid.uuid4(), email="test@test.com", hashed_password="", is_active=True)

    app.dependency_overrides[get_db_session] = override_get_db_session
    app.dependency_overrides[get_current_user] = _mock_user
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test", follow_redirects=True) as c:
        yield c
    app.dependency_overrides.clear()


# ── service-level tests ──────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_orchestrate_parses_all_pending_documents(
    session_factory, csv_bank_doc, xlsx_card_doc, open_period, monkeypatch
):
    monkeypatch.setattr(
        "app.services.orchestrate.run_orchestrator",
        AsyncMock(return_value=OrchestrationPlan(steps=[
            DocumentPlan(
                document_id=csv_bank_doc.document_id,
                resolved_type="bank_statement",
                type_reason="Bank checking activity.",
                resolved_source_account_code=100101,
                source_account_reason="Checking account on statement.",
                run_classifier=True,
            ),
            DocumentPlan(
                document_id=xlsx_card_doc.document_id,
                resolved_type="credit_card",
                type_reason="Credit card transactions.",
                resolved_source_account_code=200101,
                source_account_reason="Mastercard activity.",
                run_classifier=True,
            ),
        ])),
    )
    # Stub the classifier so we don't make real LLM calls.
    monkeypatch.setattr(
        classify_service,
        "classify_period",
        AsyncMock(return_value=0),
    )

    async with session_factory() as session:
        result = await orchestrate_service.orchestrate_parse(session, open_period.period_id)

    assert result.parsed == 2
    assert result.failed == 0
    assert result.classifier_ran is True

    async with session_factory() as session:
        rows = (await session.scalars(select(RawTransaction))).all()
        docs = (await session.scalars(select(Document))).all()
    assert len(rows) == 3  # 2 from bank, 1 from card
    assert all(d.parse_status == "complete" for d in docs)


@pytest.mark.asyncio
async def test_orchestrate_corrects_wrong_document_type(
    session_factory, mislabeled_csv_doc, open_period, monkeypatch
):
    """When the user uploads a CSV under the wrong document_type, the
    orchestrator's resolved_type should override and the document_type column
    should be updated before parse runs."""
    monkeypatch.setattr(
        "app.services.orchestrate.run_orchestrator",
        AsyncMock(return_value=OrchestrationPlan(steps=[
            DocumentPlan(
                document_id=mislabeled_csv_doc.document_id,
                resolved_type="bank_statement",  # correction
                type_reason="Header shows Balance column typical of a checking account.",
                resolved_source_account_code=100101,
                source_account_reason="Checking account.",
                run_classifier=True,
            ),
        ])),
    )
    monkeypatch.setattr(
        classify_service,
        "classify_period",
        AsyncMock(return_value=0),
    )

    async with session_factory() as session:
        result = await orchestrate_service.orchestrate_parse(session, open_period.period_id)

    assert result.parsed == 1

    async with session_factory() as session:
        doc = await session.get(Document, mislabeled_csv_doc.document_id)
    assert doc.document_type == "bank_statement"
    assert doc.parse_status == "complete"


@pytest.mark.asyncio
async def test_orchestrate_skips_classifier_when_no_bank_or_credit(
    session_factory, open_period, upload_root, monkeypatch
):
    """When the plan has only non-classifiable documents (e.g. paystub), the
    classifier should NOT be invoked."""
    period_dir = upload_root / str(open_period.period_id)
    period_dir.mkdir(parents=True, exist_ok=True)
    file_path = period_dir / "paystub.pdf"
    file_path.write_bytes(b"%PDF-1.4 placeholder")

    async with session_factory() as session:
        doc = Document(
            period_id=open_period.period_id,
            document_type="paystub",
            file_name="paystub.pdf",
            file_path=str(file_path),
            source_account_code=100101,
            parse_status="pending",
        )
        session.add(doc)
        await session.commit()
        await session.refresh(doc)

    monkeypatch.setattr(
        "app.services.orchestrate.extract_pdf_text",
        lambda path: "stub paystub text",
    )
    monkeypatch.setattr(
        "app.services.orchestrate.run_orchestrator",
        AsyncMock(return_value=OrchestrationPlan(steps=[
            DocumentPlan(
                document_id=doc.document_id,
                resolved_type="paystub",
                type_reason="Paystub format.",
                resolved_source_account_code=100101,
                source_account_reason="Net pay deposited to checking.",
                run_classifier=False,
            ),
        ])),
    )
    classifier_mock = AsyncMock(return_value=0)
    monkeypatch.setattr(classify_service, "classify_period", classifier_mock)
    # Also stub the paystub extractor — parse path will try to run it.
    from app.agents.paystub import ExtractedPaystub, ExtractedPaystubs, PaystubLine
    monkeypatch.setattr(
        "app.services.parse.extract_pdf_text", lambda path: "stub paystub text"
    )
    monkeypatch.setattr(
        "app.services.parse.run_paystub_extractor",
        AsyncMock(return_value=ExtractedPaystubs(paystubs=[
            ExtractedPaystub(
                pay_date=date(2026, 1, 15),
                lines=[PaystubLine(label="NET PAY", amount=Decimal("100"), kind="net_pay")],
                gross_pay=Decimal("100"),
                net_pay=Decimal("100"),
            )
        ])),
    )

    async with session_factory() as session:
        result = await orchestrate_service.orchestrate_parse(session, open_period.period_id)

    assert result.parsed == 1
    assert result.classifier_ran is False
    classifier_mock.assert_not_awaited()


@pytest.mark.asyncio
async def test_orchestrate_one_failure_does_not_block_others(
    session_factory, csv_bank_doc, open_period, upload_root, monkeypatch
):
    """A document the orchestrator routes to an unsupported shape (here: paystub
    with a .csv extension) should fail per-doc without stopping the rest."""
    period_dir = upload_root / str(open_period.period_id)
    bad_file = period_dir / "bad.csv"
    bad_file.write_text("garbage,here\nfoo,bar\n")

    async with session_factory() as session:
        bad = Document(
            period_id=open_period.period_id,
            document_type="bank_statement",
            file_name="bad.csv",
            file_path=str(bad_file),
            source_account_code=100101,
            parse_status="pending",
        )
        session.add(bad)
        await session.commit()
        await session.refresh(bad)

    monkeypatch.setattr(
        "app.services.orchestrate.run_orchestrator",
        AsyncMock(return_value=OrchestrationPlan(steps=[
            DocumentPlan(
                document_id=csv_bank_doc.document_id,
                resolved_type="bank_statement",
                type_reason="ok",
                resolved_source_account_code=100101,
                source_account_reason="Checking account.",
                run_classifier=False,
            ),
            DocumentPlan(
                document_id=bad.document_id,
                # Force a path that mismatches the file: paystub requires .pdf
                resolved_type="paystub",
                type_reason="(forced bad route to trigger failure)",
                resolved_source_account_code=100101,
                source_account_reason="Checking account.",
                run_classifier=False,
            ),
        ])),
    )

    async with session_factory() as session:
        result = await orchestrate_service.orchestrate_parse(session, open_period.period_id)

    assert result.parsed == 1
    assert result.failed == 1


@pytest.mark.asyncio
async def test_orchestrate_unresolved_source_account_is_needs_review(
    session_factory, csv_bank_doc, open_period, monkeypatch
):
    """When the orchestrator can't match a source account, the step should be
    reported as needs_review and the document should remain at parse_status=pending
    so the user can assign an account and click Parse on that row."""
    monkeypatch.setattr(
        "app.services.orchestrate.run_orchestrator",
        AsyncMock(return_value=OrchestrationPlan(steps=[
            DocumentPlan(
                document_id=csv_bank_doc.document_id,
                resolved_type="bank_statement",
                type_reason="ok",
                resolved_source_account_code=None,
                source_account_reason="No account name or last-4 found in content.",
                run_classifier=True,
            ),
        ])),
    )
    classifier_mock = AsyncMock(return_value=0)
    monkeypatch.setattr(classify_service, "classify_period", classifier_mock)

    async with session_factory() as session:
        result = await orchestrate_service.orchestrate_parse(session, open_period.period_id)

    assert result.parsed == 0
    assert result.failed == 0
    assert result.needs_review == 1
    assert result.classifier_ran is False
    classifier_mock.assert_not_awaited()

    async with session_factory() as session:
        doc = await session.get(Document, csv_bank_doc.document_id)
    assert doc.parse_status == "pending"


@pytest.mark.asyncio
async def test_orchestrate_no_pending_documents_is_noop(
    session_factory, open_period, monkeypatch
):
    called = AsyncMock()
    monkeypatch.setattr("app.services.orchestrate.run_orchestrator", called)

    async with session_factory() as session:
        result = await orchestrate_service.orchestrate_parse(session, open_period.period_id)

    assert result.parsed == 0
    assert result.failed == 0
    assert result.steps == []
    called.assert_not_awaited()


# ── digest scrubbing ─────────────────────────────────────────────────────────
#
# The 2500-char peek is the highest-PII payload in the pipeline: the top of a
# statement is exactly where the holder's name, address, and full account number
# sit. These tests need no DB — `_build_digest` only reads the file from disk.


def _digest_for(tmp_path, name: str, contents: str, keep_terms=()):
    path = tmp_path / name
    path.write_text(contents)
    doc = Document(
        document_id=uuid.uuid4(),
        period_id=uuid.uuid4(),
        document_type="unknown",
        file_name=name,
        file_path=str(path),
        parse_status="pending",
    )
    return orchestrate_service._build_digest(doc, keep_terms)


def test_csv_digest_peek_is_scrubbed(tmp_path):
    digest = _digest_for(
        tmp_path,
        "bank.csv",
        "Date,Description,ChkRef,Amount\n"
        "1/2/26,COFFEE SHOP 88293847192830429,,($5.00)\n",
    )
    assert "88293847192830429" not in digest.content_peek
    assert "[NUM ••0429]" in digest.content_peek
    # The merchant, date, and amount all still route the document.
    assert "COFFEE SHOP" in digest.content_peek
    assert "($5.00)" in digest.content_peek


def test_csv_digest_keeps_institution_name_via_keep_terms(tmp_path):
    digest = _digest_for(
        tmp_path,
        "bank.csv",
        "Date,Description,Amount\n1/2/26,Bank OZK Checking transfer,($5.00)\n",
        keep_terms=["Bank OZK Checking"],
    )
    assert "Bank OZK Checking" in digest.content_peek


def test_digest_file_name_is_scrubbed_but_keeps_its_last4(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "pii_identity_terms", "Anthony Scoma")
    digest = _digest_for(
        tmp_path, "Anthony Scoma Credit Card - 6120.csv", "Date,Description,Amount\n1/2/26,X,1\n"
    )
    assert "Anthony Scoma" not in digest.file_name
    # "6120" is the account last-4 the orchestrator matches on.
    assert "6120" in digest.file_name


def test_digest_is_not_scrubbed_when_the_flag_is_off(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "scrub_before_llm", False)
    digest = _digest_for(
        tmp_path, "bank.csv", "Date,Description,Amount\n1/2/26,REF 88293847192830429,1\n"
    )
    assert "88293847192830429" in digest.content_peek


NEW_FORMAT_CSV = (
    "Account Number,Post Date,Check,Description,Debit,Credit,Status,Balance\n"
    '"{acct}",7/30/2026,1042,"WAL-MART ASSOCS. PAYROLL",,2541.90,Posted,\n'
)


@pytest.mark.parametrize("acct", ["2122902212", "21229022", "2122902"])
def test_identifier_columns_are_masked_in_the_digest_at_any_length(tmp_path, acct):
    # The generic digit-run rule needs 9+ digits, so before `redact_columns` the
    # 7- and 8-digit cases reached the orchestrator prompt in full.
    digest = _digest_for(tmp_path, "bank.csv", NEW_FORMAT_CSV.format(acct=acct))
    assert acct not in digest.content_peek
    assert f"••{acct[-4:]}" in digest.content_peek
    # The last 4 is what the orchestrator matches a statement to its account on.


def test_check_number_column_is_masked_in_the_digest(tmp_path):
    digest = _digest_for(tmp_path, "bank.csv", NEW_FORMAT_CSV.format(acct="2122902212"))
    assert "Check=••1042" in digest.content_peek


def test_digest_keeps_the_columns_the_orchestrator_routes_on(tmp_path):
    digest = _digest_for(tmp_path, "bank.csv", NEW_FORMAT_CSV.format(acct="2122902212"))
    assert "WAL-MART ASSOCS. PAYROLL" in digest.content_peek
    assert "2541.90" in digest.content_peek
    assert "Post Date=7/30/2026" in digest.content_peek


def test_unreadable_file_error_is_scrubbed_but_stays_diagnostic(tmp_path):
    # The read-failure peek quotes the file name, so it carries exactly the PII the
    # rest of the digest is scrubbed for. It is redacted like anything else — while
    # staying readable enough to diagnose the failure.
    digest = _digest_for(tmp_path, "Tony Scoma Statement 1234567890.csv", "")
    assert "could not read file" in digest.content_peek
    assert "1234567890" not in digest.content_peek
    assert "1234567890" not in digest.file_name


# ── HTTP route tests ─────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_route_returns_orchestration_result(
    client: AsyncClient, csv_bank_doc, open_period, monkeypatch
):
    monkeypatch.setattr(
        "app.services.orchestrate.run_orchestrator",
        AsyncMock(return_value=OrchestrationPlan(steps=[
            DocumentPlan(
                document_id=csv_bank_doc.document_id,
                resolved_type="bank_statement",
                type_reason="ok",
                resolved_source_account_code=100101,
                source_account_reason="Checking account.",
                run_classifier=False,
            ),
        ])),
    )

    response = await client.post(
        f"/api/v1/periods/{open_period.period_id}/orchestrate-parse"
    )
    assert response.status_code == 200
    body = response.json()
    assert body["parsed"] == 1
    assert body["failed"] == 0


@pytest.mark.asyncio
async def test_route_returns_502_when_orchestrator_agent_fails(
    client: AsyncClient, csv_bank_doc, open_period, monkeypatch
):
    monkeypatch.setattr(
        "app.services.orchestrate.run_orchestrator",
        AsyncMock(side_effect=AgentError("LLM exploded")),
    )

    response = await client.post(
        f"/api/v1/periods/{open_period.period_id}/orchestrate-parse"
    )
    assert response.status_code == 502
