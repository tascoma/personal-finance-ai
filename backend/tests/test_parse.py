"""Tests for the Parse phase — service + HTTP routes.

Real LLM calls are stubbed via monkeypatch on the run_* functions. CSV/XLSX paths
exercise the deterministic mapper end-to-end with real files.
"""

from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest.mock import AsyncMock

import openpyxl
import pytest
import pytest_asyncio
from httpx import AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.agents.mortgage import ExtractedMortgage
from app.agents.paystub import ExtractedPaystub, ExtractedPaystubs, PaystubLine
from app.agents.statement import ExtractedStatement, ExtractedTxn
from app.core.config import settings
from app.databases import Base
from app.models.account import Account
from app.models.document import Document
from app.models.journal import JournalEntry, JournalLine
from app.models.raw_transaction import RawTransaction
from app.services import document as document_service
from app.services import parse as parse_service
from app.services import period as period_service
from app.services.file_readers import (
    ParseError,
    extract_csv_rows,
    extract_pdf_text,
    extract_xlsx_rows,
)

TEST_DB_URL = "sqlite+aiosqlite:///:memory:"


# ── fixtures ─────────────────────────────────────────────────────────────────


@pytest_asyncio.fixture
async def session_factory():
    eng = create_async_engine(TEST_DB_URL, echo=False)
    async with eng.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(eng, expire_on_commit=False)
    async with factory() as session:
        session.add_all(
            [
                Account(
                    account_code=100101,
                    account_name="Checking",
                    account_type="Asset",
                    sub_category="Cash",
                    normal_balance="debit",
                    is_memo=False,
                    is_active=True,
                ),
                Account(
                    account_code=200101,
                    account_name="Mastercard",
                    account_type="Liability",
                    sub_category="Credit Cards",
                    normal_balance="credit",
                    is_memo=False,
                    is_active=True,
                ),
                Account(
                    account_code=400101,
                    account_name="Salary - Regular",
                    account_type="Income",
                    sub_category="Earned",
                    normal_balance="credit",
                    paystub_mapping="REGULAR EARNING",
                    is_memo=False,
                    is_active=True,
                ),
                Account(
                    account_code=580101,
                    account_name="Health Insurance - Medical",
                    account_type="Expense",
                    sub_category="Benefits",
                    normal_balance="debit",
                    paystub_mapping="INS MED U *",
                    is_memo=False,
                    is_active=True,
                ),
                Account(
                    account_code=110102,
                    account_name="ASPP",
                    account_type="Asset",
                    sub_category="Investments",
                    normal_balance="debit",
                    paystub_mapping="CO STK CONT|STOCK PURCH",
                    is_memo=False,
                    is_active=True,
                ),
                Account(
                    account_code=300102,
                    account_name="Prior Period Net Worth",
                    account_type="Equity",
                    sub_category="Retained Equity",
                    normal_balance="credit",
                    is_memo=False,
                    is_active=True,
                ),
            ]
        )
        await session.commit()
    yield factory
    await eng.dispose()


@pytest_asyncio.fixture
async def open_period(session_factory):
    async with session_factory() as session:
        period = await period_service.create_period(session, 2026, 1)
    return period


@pytest.fixture
def upload_root(tmp_path, monkeypatch):
    root = tmp_path / "uploads"
    monkeypatch.setattr(document_service, "UPLOAD_ROOT", root)
    return root


@pytest_asyncio.fixture
async def csv_document(session_factory, open_period, upload_root):
    """A bank-statement CSV document on disk with three rows."""
    period_dir = upload_root / str(open_period.period_id)
    period_dir.mkdir(parents=True)
    file_path = period_dir / "bank.csv"
    file_path.write_text(
        "Date,Description,ChkRef,Amount,Balance\n"
        "1/2/26,COFFEE SHOP,,($5.00),$995.00\n"
        '1/3/26,DIRECT DEPOSIT,,"$2,000.00","$2,995.00"\n'
        '1/4/26,GROCERY STORE,,($45.50),"$2,949.50"\n',
    )
    async with session_factory() as session:
        doc = Document(
            period_id=open_period.period_id,
            document_type="bank_statement",
            file_name="bank.csv",
            file_path=str(file_path),
            source_account_code=100101,
            parse_status="pending",
        )
        session.add(doc)
        await session.commit()
        await session.refresh(doc)
    return doc


@pytest_asyncio.fixture
async def xlsx_document(session_factory, open_period, upload_root):
    """A credit-card XLSX document on disk with two rows."""
    period_dir = upload_root / str(open_period.period_id)
    period_dir.mkdir(parents=True, exist_ok=True)
    file_path = period_dir / "card.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Date", "Transaction", "Name", "Memo", "Amount"])
    ws.append([date(2026, 1, 5), "DEBIT", "WALMART", "", -42.10])
    ws.append([date(2026, 1, 6), "CREDIT", "REFUND", "", 17.99])
    wb.save(file_path)

    async with session_factory() as session:
        doc = Document(
            period_id=open_period.period_id,
            document_type="credit_card",
            file_name="card.xlsx",
            file_path=str(file_path),
            source_account_code=200101,
            parse_status="pending",
        )
        session.add(doc)
        await session.commit()
        await session.refresh(doc)
    return doc


@pytest_asyncio.fixture
async def paystub_document(session_factory, open_period, upload_root):
    """A paystub PDF document. The PDF text extractor is monkeypatched."""
    period_dir = upload_root / str(open_period.period_id)
    period_dir.mkdir(parents=True, exist_ok=True)
    file_path = period_dir / "paystub.pdf"
    file_path.write_bytes(b"%PDF-1.4 placeholder")

    async with session_factory() as session:
        doc = Document(
            period_id=open_period.period_id,
            document_type="paystub",
            file_name="paystub.pdf",
            file_path=str(file_path),
            source_account_code=None,
            parse_status="pending",
        )
        session.add(doc)
        await session.commit()
        await session.refresh(doc)
    return doc


@pytest_asyncio.fixture
async def statement_pdf_document(session_factory, open_period, upload_root):
    """A bank-statement PDF document. The PDF text extractor is monkeypatched."""
    period_dir = upload_root / str(open_period.period_id)
    period_dir.mkdir(parents=True, exist_ok=True)
    file_path = period_dir / "statement.pdf"
    file_path.write_bytes(b"%PDF-1.4 placeholder")

    async with session_factory() as session:
        doc = Document(
            period_id=open_period.period_id,
            document_type="bank_statement",
            file_name="statement.pdf",
            file_path=str(file_path),
            source_account_code=100101,
            parse_status="pending",
        )
        session.add(doc)
        await session.commit()
        await session.refresh(doc)
    return doc


# The identity block a real statement carries at the top of page 1.
RAW_STATEMENT_TEXT = (
    "1-234-56789-1234567-123-456-789-012-345\n"
    "ANYTOWN BANK\n"
    "1234 N MAIN ST APT#305\n"
    "SPRINGFIELD, IL 62704-1234\n"
    "Call 555-123-4567\n"
    "AccountNumber 1234567890\n"
    "Checking\n"
    "01/05 COFFEE SHOP 5.00\n"
)


# ── service-level tests ──────────────────────────────────────────────────────


async def test_csv_statement_parses_rows(session_factory, csv_document):
    async with session_factory() as session:
        count = await parse_service.parse_document(session, csv_document.document_id)
    assert count == 3

    async with session_factory() as session:
        rows = (await session.scalars(select(RawTransaction))).all()
        doc = await session.get(Document, csv_document.document_id)
    assert len(rows) == 3
    assert all(r.status == "staged" for r in rows)
    by_desc = {r.description: r for r in rows}
    assert by_desc["COFFEE SHOP"].amount == Decimal("-5.00")
    assert by_desc["DIRECT DEPOSIT"].amount == Decimal("2000.00")
    assert by_desc["GROCERY STORE"].amount == Decimal("-45.50")
    assert doc.parse_status == "complete"
    assert doc.llm_model is None  # CSV path doesn't use the LLM
    # Document.parsed_at is TIMESTAMP WITHOUT TIME ZONE — asyncpg rejects tz-aware
    # values on this column, so the service must write a naive UTC datetime.
    assert doc.parsed_at is not None
    assert doc.parsed_at.tzinfo is None


@pytest_asyncio.fixture
async def new_format_csv_document(session_factory, open_period, upload_root):
    """The bank's current export: Debit/Credit split, an account-number column,
    a Status column, and a pending row that must not be posted."""
    period_dir = upload_root / str(open_period.period_id)
    period_dir.mkdir(parents=True, exist_ok=True)
    file_path = period_dir / "AccountHistory.csv"
    file_path.write_text(
        "Account Number,Post Date,Check,Description,Debit,Credit,Status,Balance\n"
        '"2122902212",7/30/2026,,"WAL-MART ASSOCS. PAYROLL XXXXXX3311",,2541.90,Posted,\n'
        '"2122902212",7/28/2026,,"XX3178 MT DDA DEBIT VENMO  Zach John New York NY",160.00,,Posted,\n'
        '"2122902212",7/27/2026,,"PENDING COFFEE",4.50,,Pending,\n'
    )
    async with session_factory() as session:
        doc = Document(
            period_id=open_period.period_id,
            document_type="bank_statement",
            file_name="AccountHistory.csv",
            file_path=str(file_path),
            source_account_code=100101,
            parse_status="pending",
        )
        session.add(doc)
        await session.commit()
        await session.refresh(doc)
    return doc


async def test_new_bank_format_parses_split_debit_credit(
    session_factory, new_format_csv_document
):
    async with session_factory() as session:
        count = await parse_service.parse_document(
            session, new_format_csv_document.document_id
        )
    assert count == 2  # the Pending row is not posted

    async with session_factory() as session:
        rows = (await session.scalars(select(RawTransaction))).all()
        doc = await session.get(Document, new_format_csv_document.document_id)

    by_desc = {r.description: r for r in rows}
    assert by_desc["WAL-MART ASSOCS. PAYROLL XXXXXX3311"].amount == Decimal("2541.90")
    assert by_desc["XX3178 MT DDA DEBIT VENMO  Zach John New York NY"].amount == Decimal("-160.00")
    assert not any("PENDING" in d for d in by_desc)
    # Resolved deterministically — the schema-mapper agent is not consulted.
    assert doc.llm_model is None


async def test_unknown_layout_falls_back_to_the_schema_mapper(
    session_factory, open_period, upload_root, monkeypatch
):
    from app.agents.statement import ResolvedColumns
    from app.services import statement_mapper

    async def fake_schema_mapper(prompt: str) -> ResolvedColumns:
        return ResolvedColumns(
            date_column="Fecha",
            description_columns=["Concepto"],
            debit_column="Cargo",
            credit_column="Abono",
        )

    monkeypatch.setattr(statement_mapper, "run_schema_mapper", fake_schema_mapper)

    period_dir = upload_root / str(open_period.period_id)
    period_dir.mkdir(parents=True, exist_ok=True)
    file_path = period_dir / "foreign.csv"
    file_path.write_text(
        "Fecha,Concepto,Cargo,Abono\n"
        "2026-03-01,SUPERMERCADO,54.20,\n"
        "2026-03-02,NOMINA,,1200.00\n"
    )
    async with session_factory() as session:
        doc = Document(
            period_id=open_period.period_id,
            document_type="bank_statement",
            file_name="foreign.csv",
            file_path=str(file_path),
            source_account_code=100101,
            parse_status="pending",
        )
        session.add(doc)
        await session.commit()
        await session.refresh(doc)

    async with session_factory() as session:
        count = await parse_service.parse_document(session, doc.document_id)
    assert count == 2

    async with session_factory() as session:
        rows = (await session.scalars(select(RawTransaction))).all()
        parsed = await session.get(Document, doc.document_id)

    by_desc = {r.description: r.amount for r in rows}
    assert by_desc["SUPERMERCADO"] == Decimal("-54.20")
    assert by_desc["NOMINA"] == Decimal("1200.00")
    # The agent resolved the layout, so the document records which model did it.
    assert parsed.llm_model == settings.anthropic_model


async def test_xlsx_credit_card_signs_match(session_factory, xlsx_document):
    async with session_factory() as session:
        count = await parse_service.parse_document(session, xlsx_document.document_id)
    assert count == 2

    async with session_factory() as session:
        rows = (await session.scalars(select(RawTransaction).order_by(RawTransaction.amount))).all()
    assert rows[0].description == "WALMART"
    assert rows[0].amount == Decimal("-42.10")
    assert rows[1].description == "REFUND"
    assert rows[1].amount == Decimal("17.99")


async def test_paystub_maps_known_labels_and_flags_unknown(
    session_factory, paystub_document, monkeypatch
):
    monkeypatch.setattr(parse_service, "extract_pdf_text", lambda path: "stub paystub text")
    monkeypatch.setattr(
        "app.services.parse.run_paystub_extractor",
        AsyncMock(return_value=ExtractedPaystubs(
            paystubs=[
                ExtractedPaystub(
                    pay_date=date(2026, 1, 15),
                    lines=[
                        PaystubLine(label="REGULAR EARNING", amount=Decimal("3000"), kind="earning"),
                        PaystubLine(label="INS MED U PT", amount=Decimal("125"), kind="deduction"),
                        PaystubLine(label="SOMETHING WEIRD", amount=Decimal("10"), kind="deduction"),
                        PaystubLine(label="NET PAY", amount=Decimal("2865"), kind="net_pay"),
                    ],
                    gross_pay=Decimal("3000"),
                    net_pay=Decimal("2865"),
                )
            ]
        )),
    )
    async with session_factory() as session:
        count = await parse_service.parse_document(session, paystub_document.document_id)
    # net_pay is now included as a transaction line; 3 P&L + 1 net_pay = 4 rows.
    assert count == 4

    async with session_factory() as session:
        rows = (await session.scalars(select(RawTransaction))).all()
        doc = await session.get(Document, paystub_document.document_id)

    by_desc = {r.description: r for r in rows}
    assert by_desc["REGULAR EARNING"].suggested_account_code == 400101
    assert by_desc["REGULAR EARNING"].classifier_confidence == Decimal("1.000")
    assert by_desc["REGULAR EARNING"].is_flagged is False
    assert by_desc["INS MED U PT"].suggested_account_code == 580101
    assert by_desc["INS MED U PT"].amount == Decimal("-125")  # deduction → negative
    assert by_desc["SOMETHING WEIRD"].suggested_account_code is None
    assert by_desc["SOMETHING WEIRD"].is_flagged is True
    assert by_desc["SOMETHING WEIRD"].classifier_confidence == Decimal("0")
    assert doc.llm_model is not None  # paystub path used the LLM


async def test_paystub_with_two_periods_extracts_all_lines(
    session_factory, paystub_document, monkeypatch
):
    monkeypatch.setattr(parse_service, "extract_pdf_text", lambda path: "stub two-period paystub text")
    monkeypatch.setattr(
        "app.services.parse.run_paystub_extractor",
        AsyncMock(return_value=ExtractedPaystubs(
            paystubs=[
                ExtractedPaystub(
                    pay_date=date(2026, 1, 15),
                    lines=[
                        PaystubLine(label="REGULAR EARNING", amount=Decimal("3000"), kind="earning"),
                        PaystubLine(label="NET PAY", amount=Decimal("2750"), kind="net_pay"),
                    ],
                    gross_pay=Decimal("3000"),
                    net_pay=Decimal("2750"),
                ),
                ExtractedPaystub(
                    pay_date=date(2025, 12, 31),
                    lines=[
                        PaystubLine(label="REGULAR EARNING", amount=Decimal("3000"), kind="earning"),
                        PaystubLine(label="NET PAY", amount=Decimal("2750"), kind="net_pay"),
                    ],
                    gross_pay=Decimal("3000"),
                    net_pay=Decimal("2750"),
                ),
            ]
        )),
    )
    async with session_factory() as session:
        count = await parse_service.parse_document(session, paystub_document.document_id)
    # 1 earning + 1 net_pay per period × 2 periods = 4 rows
    assert count == 4

    async with session_factory() as session:
        rows = (await session.scalars(select(RawTransaction))).all()
    dates = {r.txn_date for r in rows}
    assert dates == {date(2026, 1, 15), date(2025, 12, 31)}


# ── PDF statement extraction + scrubbing ─────────────────────────────────────


async def test_pdf_statement_parses_rows_and_scrubs_the_prompt(
    session_factory, statement_pdf_document, monkeypatch
):
    """The full-raw-text path: whatever pdfplumber returns becomes the prompt."""
    monkeypatch.setattr(settings, "pii_identity_terms", "Anthony Scoma")
    monkeypatch.setattr(
        parse_service, "extract_pdf_text", lambda path: RAW_STATEMENT_TEXT + "ANTHONY SCOMA\n"
    )
    extractor = AsyncMock(return_value=ExtractedStatement(
        transactions=[
            ExtractedTxn(txn_date=date(2026, 1, 5), description="COFFEE SHOP", amount=-5.00),
        ]
    ))
    monkeypatch.setattr("app.services.parse.run_statement_extractor", extractor)

    async with session_factory() as session:
        count = await parse_service.parse_document(session, statement_pdf_document.document_id)
    assert count == 1

    prompt = extractor.await_args.args[0]
    # Nothing identifying survives.
    assert "1234567890" not in prompt
    assert "555-123-4567" not in prompt
    assert "1234 N MAIN ST" not in prompt
    assert "SPRINGFIELD, IL" not in prompt
    assert "ANTHONY SCOMA" not in prompt
    # The signals the pipeline depends on do.
    assert "[ACCT ••7890]" in prompt
    assert "Checking" in prompt
    assert "01/05 COFFEE SHOP 5.00" in prompt

    async with session_factory() as session:
        rows = (await session.scalars(select(RawTransaction))).all()
        doc = await session.get(Document, statement_pdf_document.document_id)
    # The DB keeps the extractor's verbatim description; only the prompt is scrubbed.
    assert [r.description for r in rows] == ["COFFEE SHOP"]
    assert doc.llm_model is not None


async def test_chart_of_accounts_names_survive_the_scrub(
    session_factory, statement_pdf_document, monkeypatch
):
    """`Checking` is a seeded account name — the orchestrator matches on it."""
    monkeypatch.setattr(
        parse_service,
        "extract_pdf_text",
        lambda path: "Checking 1234 N MAIN ST\nAccountNumber 1234567890\n",
    )
    extractor = AsyncMock(return_value=ExtractedStatement(transactions=[]))
    monkeypatch.setattr("app.services.parse.run_statement_extractor", extractor)

    async with session_factory() as session:
        await parse_service.parse_document(session, statement_pdf_document.document_id)

    prompt = extractor.await_args.args[0]
    assert "Checking" in prompt
    assert "1234567890" not in prompt


async def test_scrub_before_llm_false_sends_raw_text(
    session_factory, statement_pdf_document, monkeypatch
):
    monkeypatch.setattr(settings, "scrub_before_llm", False)
    monkeypatch.setattr(parse_service, "extract_pdf_text", lambda path: RAW_STATEMENT_TEXT)
    extractor = AsyncMock(return_value=ExtractedStatement(transactions=[]))
    monkeypatch.setattr("app.services.parse.run_statement_extractor", extractor)

    async with session_factory() as session:
        await parse_service.parse_document(session, statement_pdf_document.document_id)

    assert extractor.await_args.args[0] == RAW_STATEMENT_TEXT


async def test_paystub_prompt_is_scrubbed(session_factory, paystub_document, monkeypatch):
    monkeypatch.setattr(
        parse_service, "extract_pdf_text", lambda path: "Emp #: 123456789\nREGULAR EARNING 3,000.00\n"
    )
    extractor = AsyncMock(return_value=ExtractedPaystubs(paystubs=[]))
    monkeypatch.setattr("app.services.parse.run_paystub_extractor", extractor)

    async with session_factory() as session:
        await parse_service.parse_document(session, paystub_document.document_id)

    prompt = extractor.await_args.args[0]
    assert "123456789" not in prompt
    # Payroll labels and amounts are the extraction payload — never touched.
    assert "REGULAR EARNING 3,000.00" in prompt


async def test_account_keep_terms_returns_active_account_names(session_factory):
    async with session_factory() as session:
        terms = await parse_service.account_keep_terms(session)
    assert "Checking" in terms
    assert "Mastercard" in terms


async def test_reparse_replaces_prior_rows(session_factory, csv_document):
    async with session_factory() as session:
        await parse_service.parse_document(session, csv_document.document_id)

    async with session_factory() as session:
        await parse_service.parse_document(session, csv_document.document_id)

    async with session_factory() as session:
        rows = (await session.scalars(select(RawTransaction))).all()
    assert len(rows) == 3  # not 6


async def test_duplicate_documents_flagged_on_second_pass(
    session_factory, csv_document, open_period, upload_root
):
    async with session_factory() as session:
        await parse_service.parse_document(session, csv_document.document_id)

    # Upload the same CSV under a different document.
    period_dir = upload_root / str(open_period.period_id)
    file_path = period_dir / "bank-copy.csv"
    file_path.write_text(Path(csv_document.file_path).read_text())

    async with session_factory() as session:
        dup = Document(
            period_id=open_period.period_id,
            document_type="bank_statement",
            file_name="bank-copy.csv",
            file_path=str(file_path),
            source_account_code=100101,
            parse_status="pending",
        )
        session.add(dup)
        await session.commit()
        await session.refresh(dup)

    async with session_factory() as session:
        await parse_service.parse_document(session, dup.document_id)

    async with session_factory() as session:
        dup_rows = (await session.scalars(
            select(RawTransaction).where(RawTransaction.document_id == dup.document_id)
        )).all()
    assert len(dup_rows) == 3
    assert all(r.is_duplicate for r in dup_rows)


async def test_parse_refuses_unknown_doc_type(
    session_factory, open_period, upload_root
):
    """An 'unknown' doc_type must not silently fall through to the statement extractor."""
    period_dir = upload_root / str(open_period.period_id)
    period_dir.mkdir(parents=True, exist_ok=True)
    file_path = period_dir / "mystery.pdf"
    file_path.write_bytes(b"%PDF-1.4 placeholder")

    async with session_factory() as session:
        doc = Document(
            period_id=open_period.period_id,
            document_type="unknown",
            file_name="mystery.pdf",
            file_path=str(file_path),
            source_account_code=None,
            parse_status="pending",
        )
        session.add(doc)
        await session.commit()
        await session.refresh(doc)

    async with session_factory() as session:
        with pytest.raises(ParseError, match="unknown"):
            await parse_service.parse_document(session, doc.document_id)

        refreshed = await session.get(Document, doc.document_id)
        assert refreshed.parse_status == "failed"


async def test_parse_blocked_when_period_not_open(
    session_factory, csv_document, open_period
):
    async with session_factory() as session:
        await period_service.update_status(
            session, open_period.period_id, "pending_review"
        )

    async with session_factory() as session:
        with pytest.raises(ParseError):
            await parse_service.parse_document(session, csv_document.document_id)


async def test_parse_period_continues_past_failures(
    session_factory, csv_document, open_period, upload_root
):
    period_dir = upload_root / str(open_period.period_id)
    bad_path = period_dir / "broken.csv"
    bad_path.write_text("not,a,real,statement\nfoo,bar,baz,qux\n")

    async with session_factory() as session:
        bad = Document(
            period_id=open_period.period_id,
            document_type="bank_statement",
            file_name="broken.csv",
            file_path=str(bad_path),
            source_account_code=100101,
            parse_status="pending",
        )
        session.add(bad)
        await session.commit()
        await session.refresh(bad)

        results = await parse_service.parse_period(session, open_period.period_id)

    assert results[csv_document.document_id] == 3
    assert isinstance(results[bad.document_id], str)


# ── HTTP route tests ─────────────────────────────────────────────────────────


async def test_parse_route_returns_document_with_complete_status(
    client: AsyncClient, csv_document, open_period
):
    response = await client.post(
        f"/api/v1/periods/{open_period.period_id}/documents/{csv_document.document_id}/parse"
    )
    assert response.status_code == 200
    assert response.json()["parse_status"] == "complete"


async def test_transactions_list_rows(
    client: AsyncClient, csv_document, open_period
):
    await client.post(
        f"/api/v1/periods/{open_period.period_id}/documents/{csv_document.document_id}/parse"
    )
    response = await client.get(f"/api/v1/periods/{open_period.period_id}/transactions")
    assert response.status_code == 200
    descriptions = [t["description"] for t in response.json()]
    assert any("COFFEE SHOP" in d for d in descriptions)
    assert any("DIRECT DEPOSIT" in d for d in descriptions)


async def test_parse_route_blocked_for_non_open_period(
    client: AsyncClient, csv_document, open_period, session_factory
):
    async with session_factory() as session:
        await period_service.update_status(
            session, open_period.period_id, "pending_review"
        )
    response = await client.post(
        f"/api/v1/periods/{open_period.period_id}/documents/{csv_document.document_id}/parse"
    )
    assert response.status_code == 400
    assert "open" in response.json()["detail"].lower()


# ── mortgage statement tests ──────────────────────────────────────────────────


@pytest_asyncio.fixture
async def mortgage_document(session_factory, open_period, upload_root):
    """A mortgage statement PDF document. The PDF text extractor is monkeypatched."""
    period_dir = upload_root / str(open_period.period_id)
    period_dir.mkdir(parents=True, exist_ok=True)
    file_path = period_dir / "mortgage.pdf"
    file_path.write_bytes(b"%PDF-1.4 placeholder")

    async with session_factory() as session:
        doc = Document(
            period_id=open_period.period_id,
            document_type="mortgage_statement",
            file_name="mortgage.pdf",
            file_path=str(file_path),
            source_account_code=100101,
            parse_status="pending",
        )
        session.add(doc)
        await session.commit()
        await session.refresh(doc)
    return doc


async def test_mortgage_statement_monthly_payment(
    session_factory, mortgage_document, monkeypatch
):
    """Regular monthly statement: principal, interest, and an escrow deposit."""
    monkeypatch.setattr(
        parse_service, "extract_pdf_text", lambda path: "stub mortgage text"
    )
    monkeypatch.setattr(
        "app.services.parse.run_mortgage_extractor",
        AsyncMock(return_value=ExtractedMortgage(
            payment_date=date(2026, 1, 1),
            principal=Decimal("450.00"),
            interest=Decimal("1200.00"),
            escrow=Decimal("350.00"),
            property_tax=Decimal("0"),
            home_insurance=Decimal("0"),
        )),
    )
    async with session_factory() as session:
        count = await parse_service.parse_document(session, mortgage_document.document_id)
    assert count == 3

    async with session_factory() as session:
        rows = (await session.scalars(select(RawTransaction))).all()
        doc = await session.get(Document, mortgage_document.document_id)

    assert len(rows) == 3
    by_desc = {r.description: r for r in rows}
    assert by_desc["Mortgage Principal"].amount == Decimal("-450.00")
    assert by_desc["Mortgage Interest"].amount == Decimal("-1200.00")
    assert by_desc["Escrow Deposit"].amount == Decimal("-350.00")
    assert all(r.status == "staged" for r in rows)
    assert by_desc["Mortgage Principal"].suggested_account_code == 210102
    assert by_desc["Mortgage Interest"].suggested_account_code == 510101
    assert by_desc["Escrow Deposit"].suggested_account_code == 100202
    assert all(r.classifier_confidence == Decimal("1.000") for r in rows)
    assert doc.parse_status == "complete"
    assert doc.llm_model is not None


async def test_mortgage_escrow_disbursement(
    session_factory, mortgage_document, monkeypatch
):
    """Escrow analysis statement: all five components present."""
    monkeypatch.setattr(
        parse_service, "extract_pdf_text", lambda path: "stub mortgage text"
    )
    monkeypatch.setattr(
        "app.services.parse.run_mortgage_extractor",
        AsyncMock(return_value=ExtractedMortgage(
            payment_date=date(2026, 1, 1),
            principal=Decimal("450.00"),
            interest=Decimal("1200.00"),
            escrow=Decimal("350.00"),
            property_tax=Decimal("2100.00"),
            home_insurance=Decimal("900.00"),
        )),
    )
    async with session_factory() as session:
        count = await parse_service.parse_document(session, mortgage_document.document_id)
    assert count == 5

    async with session_factory() as session:
        rows = (await session.scalars(select(RawTransaction))).all()

    by_desc = {r.description: r for r in rows}
    assert by_desc["Escrow Deposit"].suggested_account_code == 100202
    assert by_desc["Property Tax"].suggested_account_code == 510102
    assert by_desc["Home Insurance"].suggested_account_code == 510103


async def test_mortgage_statement_omits_zero_components(
    session_factory, mortgage_document, monkeypatch
):
    """Components with zero amount are skipped (e.g. no escrow on some loans)."""
    monkeypatch.setattr(
        parse_service, "extract_pdf_text", lambda path: "stub mortgage text"
    )
    monkeypatch.setattr(
        "app.services.parse.run_mortgage_extractor",
        AsyncMock(return_value=ExtractedMortgage(
            payment_date=date(2026, 1, 1),
            principal=Decimal("500.00"),
            interest=Decimal("1100.00"),
            escrow=Decimal("0"),
            property_tax=Decimal("0"),
            home_insurance=Decimal("0"),
        )),
    )
    async with session_factory() as session:
        count = await parse_service.parse_document(session, mortgage_document.document_id)
    assert count == 2

    async with session_factory() as session:
        rows = (await session.scalars(select(RawTransaction))).all()
    descriptions = {r.description for r in rows}
    assert descriptions == {"Mortgage Principal", "Mortgage Interest"}


async def test_mortgage_must_be_pdf(
    session_factory, open_period, upload_root, monkeypatch
):
    period_dir = upload_root / str(open_period.period_id)
    period_dir.mkdir(parents=True, exist_ok=True)
    file_path = period_dir / "mortgage.csv"
    file_path.write_text("not,a,pdf\n")

    async with session_factory() as session:
        doc = Document(
            period_id=open_period.period_id,
            document_type="mortgage_statement",
            file_name="mortgage.csv",
            file_path=str(file_path),
            source_account_code=100101,
            parse_status="pending",
        )
        session.add(doc)
        await session.commit()
        await session.refresh(doc)

    async with session_factory() as session:
        with pytest.raises(ParseError, match="must be .pdf"):
            await parse_service.parse_document(session, doc.document_id)


# ── opening-balance tests ────────────────────────────────────────────────────


def _write_opening_balances_xlsx(path: Path, rows: list[tuple[int, float]]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["account_code", "balance"])
    for code, bal in rows:
        ws.append([code, bal])
    wb.save(path)


@pytest_asyncio.fixture
async def opening_balances_document(session_factory, open_period, upload_root):
    period_dir = upload_root / str(open_period.period_id)
    period_dir.mkdir(parents=True, exist_ok=True)
    file_path = period_dir / "opening.xlsx"
    _write_opening_balances_xlsx(
        file_path,
        [
            (100101, 5200.00),    # Checking — Asset, debit-normal, positive
            (110102, 18000.00),   # ASPP — Asset
            (200101, -1400.00),   # Mastercard — Liability, credit-normal, negative
        ],
    )

    async with session_factory() as session:
        doc = Document(
            period_id=open_period.period_id,
            document_type="opening_balances",
            file_name="opening.xlsx",
            file_path=str(file_path),
            source_account_code=None,
            parse_status="pending",
        )
        session.add(doc)
        await session.commit()
        await session.refresh(doc)
    return doc


async def test_opening_balances_creates_balanced_journal_entry(
    session_factory, opening_balances_document
):
    async with session_factory() as session:
        line_count = await parse_service.parse_document(session, opening_balances_document.document_id)

    # 3 input rows + 1 equity offset
    assert line_count == 4

    async with session_factory() as session:
        # No RawTransactions written
        raw = (await session.scalars(select(RawTransaction))).all()
        assert raw == []

        entries = (await session.scalars(select(JournalEntry))).all()
        assert len(entries) == 1
        entry = entries[0]
        assert entry.source_type == "adjusting"
        assert entry.is_adjusting is True
        assert entry.created_by == "python"
        assert entry.source_document_id == opening_balances_document.document_id

        lines = (await session.scalars(
            select(JournalLine).where(JournalLine.entry_id == entry.entry_id)
        )).all()

    by_account = {line.account_code: line for line in lines}
    assert by_account[100101].debit_amount == Decimal("5200.00")
    assert by_account[100101].credit_amount == Decimal("0")
    assert by_account[110102].debit_amount == Decimal("18000.00")
    assert by_account[200101].credit_amount == Decimal("1400.00")
    assert by_account[200101].debit_amount == Decimal("0")
    # Equity offset balances the entry: debits 23200, credits 1400 → 21800 credit
    assert by_account[300102].credit_amount == Decimal("21800.00")

    total_debits = sum(line.debit_amount for line in lines)
    total_credits = sum(line.credit_amount for line in lines)
    assert total_debits == total_credits


async def test_opening_balances_rejects_wrong_sign(
    session_factory, open_period, upload_root
):
    period_dir = upload_root / str(open_period.period_id)
    period_dir.mkdir(parents=True, exist_ok=True)
    file_path = period_dir / "bad-signs.xlsx"
    # Liability given as a positive number — should be negative.
    _write_opening_balances_xlsx(file_path, [(200101, 1500.00)])

    async with session_factory() as session:
        doc = Document(
            period_id=open_period.period_id,
            document_type="opening_balances",
            file_name="bad-signs.xlsx",
            file_path=str(file_path),
            parse_status="pending",
        )
        session.add(doc)
        await session.commit()
        await session.refresh(doc)

    async with session_factory() as session:
        with pytest.raises(ParseError, match="negative number"):
            await parse_service.parse_document(session, doc.document_id)


async def test_opening_balances_skips_zero_rows(
    session_factory, open_period, upload_root
):
    period_dir = upload_root / str(open_period.period_id)
    period_dir.mkdir(parents=True, exist_ok=True)
    file_path = period_dir / "zeros.xlsx"
    _write_opening_balances_xlsx(
        file_path,
        [
            (100101, 1000.00),
            (110102, 0),  # skip
            (200101, -500.00),
        ],
    )

    async with session_factory() as session:
        doc = Document(
            period_id=open_period.period_id,
            document_type="opening_balances",
            file_name="zeros.xlsx",
            file_path=str(file_path),
            parse_status="pending",
        )
        session.add(doc)
        await session.commit()
        await session.refresh(doc)

    async with session_factory() as session:
        line_count = await parse_service.parse_document(session, doc.document_id)
    # 2 non-zero + 1 offset
    assert line_count == 3


# ── file readers: encoding + error-message hygiene ───────────────────────────


def test_cp1252_csv_is_read_rather_than_rejected(tmp_path):
    # Banks still ship Windows-encoded exports; the 0x92 curly apostrophe in
    # "MCDONALD'S" is not valid UTF-8 and used to fail the whole document.
    path = tmp_path / "windows.csv"
    path.write_bytes(
        b"Date,Description,Amount\n1/2/26,MCDONALD\x92S CAF\xc9,-5.00\n"
    )
    rows = extract_csv_rows(path)
    assert rows[0]["Description"] == "MCDONALD’S CAFÉ"


def test_utf8_csv_still_wins_over_the_cp1252_fallback(tmp_path):
    path = tmp_path / "utf8.csv"
    path.write_text("Date,Description,Amount\n1/2/26,CAFÉ MØLLER,-5.00\n", encoding="utf-8")
    assert extract_csv_rows(path)[0]["Description"] == "CAFÉ MØLLER"


@pytest.mark.parametrize(
    "name,contents",
    [
        ("Tony Scoma Statement 1234567890.csv", ""),
        ("Tony Scoma Statement 1234567890.pdf", "not a pdf"),
        ("Tony Scoma Statement 1234567890.xlsx", "not a workbook"),
    ],
)
def test_read_errors_never_quote_the_upload_path(tmp_path, name, contents):
    # The message reaches the orchestrator prompt via the digest peek, and the
    # upload path carries the period UUID as well as the file name.
    path = tmp_path / name
    path.write_text(contents)
    reader = {
        ".csv": extract_csv_rows,
        ".pdf": extract_pdf_text,
        ".xlsx": extract_xlsx_rows,
    }[path.suffix]

    with pytest.raises(ParseError) as exc:
        reader(path)
    assert str(tmp_path) not in str(exc.value)
